#!/usr/bin/env python3
"""Erzeugt orte.js (Vorlage der Partnerorte) aus der Recherche-Arbeitsmappe.

Aufruf:  python3 -I scripts/orte_aus_excel.py quellen/Freizeitticket_2026-27.xlsx > orte.js
Braucht openpyxl (pip install openpyxl).

Quelle: Blatt „Partnerliste“ (Orte, Gemeinde, Kategorie, Nutzungsregeln, Preisstand) und
„Reguläre Preise“ (Einzeltarife je Partner). Übernommen werden nur Einzeleintritte/Tageskarten –
Saison-, Jahres-, Block-, Gruppen- und Mehrtageskarten fallen weg, weil jeder Besuch einzeln zählt.
Preise sind Recherche-Stand (ohne Gewähr) und bleiben in der App editierbar.
"""
import json
import re
import sys
import unicodedata

import openpyxl

# Bestehende IDs (aus der ersten Vorlage) beibehalten – Besuche verweisen darauf.
ID_FEST = {
    'Schlick 2000': 'schlick',
    'Serlesbahn Mieders': 'serlesbahnen',
    'Tivoli Freibad Innsbruck': 'tivoli',
    'Schwimmbad Hall': 'parkbad-hall',
    'Olympiaworld Innsbruck': 'olympiaworld-eis',
    'Eislaufplatz Igls': 'eislauf-igls',
    'Tirol Panorama mit Kaiserjägermuseum': 'tirol-panorama',
    'Tiroler Volkskunstmuseum': 'volkskunstmuseum',
}

# Partnername im Blatt „Reguläre Preise“ → Name(n) in der Partnerliste
TARIF_ZU_PARTNER = {
    'Kühtai–Hochoetz': ['Kühtai'],
    'Bergbahnen Rosshütte Seefeld': ['Rosshütte Seefeld'],
    'Silvapark Galtür': ['Silvretta Galtür'],
    'Kappl–See': ['Kappl'],
    'Kellerjochbahn': ['Kellerjoch'],
    'Silvretta Arena Ischgl/Samnaun': ['Ischgl (Silvretta Seilbahn)'],
    'Birkenlift & Geigenbühel': ['Seefeld Birkenlift', 'Skilift Geigenbühel'],
    'Tiroler Landesmuseen (Hofkirche, Volkskunstmuseum, Zeughaus, Tirol Panorama)': [
        'Hofkirche', 'Tirol Panorama mit Kaiserjägermuseum', 'Tiroler Volkskunstmuseum', 'Zeughaus'],
    'Tiroler Landesmuseen / Tirol Panorama / Hofkirche / Volkskunstmuseum / Zeughaus': [
        'Hofkirche', 'Tirol Panorama mit Kaiserjägermuseum', 'Tiroler Volkskunstmuseum', 'Zeughaus'],
    'Museum Goldenes Dachl / Stadtmuseum': ['Museum Goldenes Dachl'],
    'Badesee Weißlahn': ['Badesee Weißlan'],
}

AUSSCHLIESSEN = re.compile(
    r'saisonkarte|saison-karte|jahres|block|dutzend|monat|wertkarte|gruppe|führung|kabine|familie|family|aufzahlung'
    r'|card|\d\s*tage|^\d+ tage|behinderten|schul|ermäßigt|senior|schneemann|kombi|golden roof',
    re.I)

KATEGORIE = [
    (re.compile(r'ski|lift', re.I), 'bergbahn'),
    (re.compile(r'bad|see|schwimm', re.I), 'bad'),
    (re.compile(r'eislauf', re.I), 'eislauf'),
    (re.compile(r'kultur|museum', re.I), 'museum'),
]

ALLGEMEIN = ('Reguläre Nutzung', 'Eintritt laut Saison', 'Freier Eintritt laut offizieller')


def slug(s):
    s = s.lower().replace('ä', 'ae').replace('ö', 'oe').replace('ü', 'ue').replace('ß', 'ss')
    s = unicodedata.normalize('NFD', s)
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-')[:60]


def preis(v):
    """Zahl oder '€ 50,50–66,50' → erste Zahl; sonst None."""
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    m = re.search(r'(\d+(?:[.,]\d+)?)', str(v or ''))
    return round(float(m.group(1).replace(',', '.')), 2) if m else None


def ticketname(n):
    n = re.sub(r'\s+', ' ', str(n)).strip()
    n = n.replace('Berg-/Talfahrt', 'Berg- und Talfahrt')
    return n[:80]


def main(pfad):
    wb = openpyxl.load_workbook(pfad, data_only=True)

    partner = {}
    for r in wb['Partnerliste'].iter_rows(min_row=4, values_only=True):
        kat_text, name, gemeinde = r[0], r[1], r[2]
        if not name or str(name).startswith('13. Eislaufplatz'):
            continue
        name = str(name).strip()
        kat = next((k for rx, k in KATEGORIE if rx.search(str(kat_text))), 'sonstiges')
        regel = str(r[5] or '').strip()
        hinweise = []
        if regel and not regel.startswith(ALLGEMEIN):
            hinweise.append(regel.rstrip('.') + '.')
        if 'unter 19 frei' in str(r[16] or ''):
            hinweise.append('Kinder und Jugendliche unter 19 ohnehin frei (Kinderpreis 0 €).')
        if r[14]:
            hinweise.append('Preisstand: ' + str(r[14]).strip().rstrip('.') + '.')
        link = str(r[7] or '').strip()
        partner[name] = {
            'id': ID_FEST.get(name, slug(name)),
            'name': name,
            'gemeinde': str(gemeinde or '').strip(),
            'kat': kat,
            'tickets': [],
            'hinweis': ' '.join(hinweise)[:300],
            'link': link if link.startswith('https://www.freizeitticket.at/') else '',
            '_ersatz': (preis(r[11]), preis(r[12]), ticketname(r[13] or 'Eintritt')),
        }

    # Tarife einsammeln
    roh = {}
    for r in wb['Reguläre Preise'].iter_rows(min_row=4, values_only=True):
        p, t, erw, jug, kind = r[0], r[2], r[4], r[6], r[7]
        if not p or not t or AUSSCHLIESSEN.search(str(t)):
            continue
        p = str(p).strip()
        ziele = TARIF_ZU_PARTNER.get(p, [p])
        for z in ziele:
            if z in partner:
                roh.setdefault(z, []).append((ticketname(t), preis(erw), preis(kind), preis(jug)))

    for name, o in partner.items():
        zeilen = roh.get(name, [])
        nur_kind = [z for z in zeilen if z[1] is None and z[2] is not None]
        mit_erw = [z for z in zeilen if z[1] is not None]
        kinderpreis = nur_kind[0][2] if nur_kind else None
        kein_kindertarif = all(z[2] is None for z in mit_erw)
        tickets = []
        for n, e, k, j in mit_erw:
            if k is None:
                if kinderpreis is not None:
                    k = kinderpreis               # eigenes Kinderticket (z. B. Swarovski)
                elif kein_kindertarif and o['_ersatz'][1] is not None:
                    k = o['_ersatz'][1]           # Partnerliste (Landesmuseen: 0 = unter 19 frei)
                elif kein_kindertarif and j is not None:
                    k = j                         # nur ermäßigter Tarif bekannt
                else:
                    continue
            tickets.append([n, e, k])
        # gleicher Name, verschiedene Preise → nach Saison benennen
        namen = {}
        for t in tickets:
            namen.setdefault(t[0], []).append(t)
        eindeutig = []
        for n, gruppe in namen.items():
            gruppe = sorted({(t[1], t[2]) for t in gruppe})
            if len(gruppe) == 1:
                eindeutig.append([n, *gruppe[0]])
            else:
                stufen = ['Vorsaison', 'Hauptsaison', 'Topsaison'] if len(gruppe) == 3 else ['Nebensaison', 'Hauptsaison']
                for s, (e, k) in zip(stufen, gruppe):
                    eindeutig.append([f'{n} – {s}', e, k])
        if not eindeutig:
            e, k, n = o['_ersatz']
            if AUSSCHLIESSEN.search(n):
                n = 'Eintritt'
            if e is None and k is None:
                o['hinweis'] = ('Kein Preis gefunden – bitte selbst eintragen. ' + o['hinweis']).strip()[:300]
            elif k is None:
                o['hinweis'] = ('Kinderpreis nicht gefunden – bitte eintragen. ' + o['hinweis']).strip()[:300]
            eindeutig.append([n, e or 0, k or 0])
        gesehen = set()
        preise = set()
        for n, e, k in eindeutig:
            if (e, k) in preise:
                continue                          # gleicher Tarif unter anderem Namen (Doppelerfassung)
            preise.add((e, k))
            tid = slug(n) or 'ticket'
            while tid in gesehen:
                tid += '-2'
            gesehen.add(tid)
            o['tickets'].append({'id': tid, 'name': n, 'erw': e, 'kind': k})
        del o['_ersatz']

    orte = sorted(partner.values(), key=lambda o: (o['kat'], o['name']))
    ids = [o['id'] for o in orte]
    assert len(ids) == len(set(ids)), 'IDs nicht eindeutig'

    print('/* ERZEUGT von scripts/orte_aus_excel.py aus ' + pfad.split('/')[-1] + ' – nicht von Hand ändern.')
    print(' * Preise: Recherche-Stand ohne Gewähr, in der App editierbar. */')
    print('(function (root) {')
    print("  'use strict';")
    print('  const ORTE = ' + json.dumps(orte, ensure_ascii=False, indent=2).replace('\n', '\n  ') + ';')
    print('  root.FZ_ORTE_VORLAGE = ORTE;')
    print("  if (typeof module !== 'undefined' && module.exports) module.exports = ORTE;")
    print("})(typeof globalThis !== 'undefined' ? globalThis : this);")


if __name__ == '__main__':
    main(sys.argv[1])
